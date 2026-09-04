from pathlib import Path
from datetime import datetime

import pandas as pd
import streamlit as st

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


st.set_page_config(
    page_title="Wasserstein Similarity Study",
    page_icon="📊",
    layout="wide",
)

# Reduce this value if the question page is too tall for the screen.
IMAGE_HEIGHT_PX = 185

# Add your explanatory Wasserstein example diagram at this location.
EXAMPLE_DIAGRAM_PATH = "images/example/wasserstein_example.png"

# Enter the values produced by your example-plot generation script here.
# These are example values. Replace them with the actual W(A, Q) and W(B, Q)
# values printed by create_wasserstein_example.py.
EXAMPLE_DISTANCE_A_TO_Q = 0.3000
EXAMPLE_DISTANCE_B_TO_Q = 2.9597


st.markdown(
    f"""
    <style>
        #MainMenu, header, footer {{
            visibility: hidden;
        }}

        .block-container {{
            max-width: 900px;
            padding-top: 0.4rem;
            padding-bottom: 0.4rem;
        }}

        h1 {{
            font-size: 1.8rem;
            margin-top: 0.2rem;
            margin-bottom: 0.5rem;
        }}

        h2 {{
            font-size: 1.35rem;
            margin-top: 0.2rem;
            margin-bottom: 0.5rem;
        }}

        h3 {{
            font-size: 1.05rem;
            margin-top: 0.5rem;
            margin-bottom: 0.2rem;
        }}

        p {{
            margin-bottom: 0.35rem;
        }}

        /*
        This class is added only around quiz images.
        It keeps the Option A, Question Q, and Option B images
        at a consistent height on the actual quiz page.
        */
        .quiz-image [data-testid="stImage"] {{
            margin-bottom: -0.35rem;
        }}

        .quiz-image [data-testid="stImage"] img {{
            height: {IMAGE_HEIGHT_PX}px !important;
            width: 100% !important;
            object-fit: contain !important;
            border: 2px solid transparent;
            border-radius: 8px;
            transition: border-color 0.15s ease;
        }}

        .quiz-image [data-testid="stImage"] img:hover {{
            border-color: #1f77b4;
            cursor: pointer;
        }}

        /*
        The intro example diagram must not use IMAGE_HEIGHT_PX.
        It preserves its full natural aspect ratio, so the Option B
        row is not clipped or compressed.
        */
        .example-diagram [data-testid="stImage"] {{
            margin-bottom: 0.3rem;
        }}

        .example-diagram [data-testid="stImage"] img {{
            height: auto !important;
            max-height: none !important;
            width: 100% !important;
            object-fit: contain !important;
            border: none !important;
            border-radius: 0 !important;
            cursor: default !important;
        }}

        div[data-testid="stButton"] > button {{
            min-height: 42px;
            margin-top: 0.15rem;
            margin-bottom: 0.15rem;
        }}

        [data-testid="stProgress"] {{
            margin-top: 0;
            margin-bottom: 0.1rem;
        }}

        div[data-testid="stAlert"] {{
            padding-top: 0.25rem;
            padding-bottom: 0.25rem;
            margin-top: 0.15rem;
            margin-bottom: 0.15rem;
        }}

        .diagram-placeholder {{
            border: 2px dashed #8c8c8c;
            border-radius: 10px;
            padding: 3rem 1rem;
            margin: 1rem 0;
            text-align: center;
            color: #555555;
            background-color: #f7f7f7;
        }}

        .slide-note {{
            border-left: 4px solid #1f77b4;
            border-radius: 4px;
            padding: 0.7rem 0.9rem;
            margin-top: 0.8rem;
            background-color: #f0f6fc;
        }}
    </style>
    """,
    unsafe_allow_html=True,
)


QUESTION_TEXT = "Which option appears more similar to the given distribution?"

# Fill in the correct options

CORRECT_OPTIONS = [
    "B",  # Q1
    "A",  # Q2
    "B",  # Q3
    "B",  # Q4
    "A",  # Q5
    "B",  # Q6
    "B",  # Q7
    "B",  # Q8
    "B",  # Q9
    "A",  # Q10
    "B",  # Q11
    "B",  # Q12
    "A",  # Q13
    "A",  # Q14
    "B",  # Q15
    "A",  # Q16
    "B",  # Q17
    "B",  # Q18
    "A",  # Q19
    "B",  # Q20
]

# Fill these values with your actual Wasserstein distances.
#
# Each tuple represents:
# (Wasserstein distance from Option A to Question,
#  Wasserstein distance from Option B to Question)

DISTANCES = [
    (8.000, 2.279),  # Q1
    (0.111, 8.000),  # Q2
    (6.000, 0.200),  # Q3
    (6.000, 0.422),  # Q4
    (2.625, 4.000),  # Q5
    (2.663, 0.147),  # Q6
    (4.064, 2.045),  # Q7
    (2.327, 0.459),  # Q8
    (2.910, 0.640),  # Q9
    (0.422, 0.893),  # Q10
    (0.935, 0.453),  # Q11
    (0.458, 0.045),  # Q12
    (0.309, 1.193),  # Q13
    (2.006, 2.467),  # Q14
    (2.540, 1.063),  # Q15
    (0.375, 0.507),  # Q16
    (0.530, 0.458),  # Q17
    (2.830, 1.921),  # Q18
    (1.921, 2.000),  # Q19
    (2.527, 2.167),  # Q20
]



if len(CORRECT_OPTIONS) != 20:
    raise ValueError("CORRECT_OPTIONS must contain exactly 20 entries.")


if len(DISTANCES) != 20:
    raise ValueError("DISTANCES must contain exactly 20 distance pairs.")


QUESTIONS = [
    {
        "target_image": f"images/Q{i}/t.png",
        "option_a_image": f"images/Q{i}/a.png",
        "option_b_image": f"images/Q{i}/b.png",
        "correct_option": CORRECT_OPTIONS[i - 1],
        "distance_option_a": DISTANCES[i - 1][0],
        "distance_option_b": DISTANCES[i - 1][1],
    }
    for i in range(1, 21)
]


RESULTS_FILE = "results.xlsx"


def initialize_state():
    if "screen" not in st.session_state:
        st.session_state.screen = "study"

    if "question_index" not in st.session_state:
        st.session_state.question_index = 0

    if "answers" not in st.session_state:
        st.session_state.answers = []

    if "selected_option" not in st.session_state:
        st.session_state.selected_option = None

    if "quiz_finished" not in st.session_state:
        st.session_state.quiz_finished = False


def go_to_instructions():
    st.session_state.screen = "instructions"


def start_quiz():
    st.session_state.screen = "quiz"


def style_excel_sheet(worksheet, table_name):
    """Apply formatting and create an Excel table in one worksheet."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32

    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter

        maximum_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )

        worksheet.column_dimensions[column_letter].width = min(
            max(maximum_length + 3, 12),
            48,
        )

    if worksheet.max_row >= 2:
        table = Table(
            displayName=table_name,
            ref=worksheet.dimensions,
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = table_style
        worksheet.add_table(table)


def record_result():
    """Save one completed quiz attempt into a formatted Excel workbook."""

    score = sum(answer["is_correct"] for answer in st.session_state.answers)
    total_questions = len(QUESTIONS)
    percentage = 100 * score / total_questions
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    new_attempt_df = pd.DataFrame(
        [
            {
                "Timestamp": timestamp,
                "Score": score,
                "Total Questions": total_questions,
                "Percentage (%)": round(percentage, 2),
            }
        ]
    )

    new_question_details_df = pd.DataFrame(
        [
            {
                "Timestamp": timestamp,
                "Question": f"Q{answer['question']}",
                "Selected Option": f"Option {answer['selected_option']}",
                "Correct Option": f"Option {answer['correct_option']}",
                "Result": "Correct" if answer["is_correct"] else "Incorrect",
                "W(A, Q)": round(answer["distance_option_a"], 6),
                "W(B, Q)": round(answer["distance_option_b"], 6),
            }
            for answer in st.session_state.answers
        ]
    )

    results_path = Path(RESULTS_FILE)

    if results_path.exists():
        old_attempts_df = pd.read_excel(
            RESULTS_FILE,
            sheet_name="Attempts",
        )

        old_question_details_df = pd.read_excel(
            RESULTS_FILE,
            sheet_name="Question Details",
        )

        attempts_df = pd.concat(
            [old_attempts_df, new_attempt_df],
            ignore_index=True,
        )

        question_details_df = pd.concat(
            [old_question_details_df, new_question_details_df],
            ignore_index=True,
        )

    else:
        attempts_df = new_attempt_df
        question_details_df = new_question_details_df

    with pd.ExcelWriter(
        RESULTS_FILE,
        engine="openpyxl",
        mode="w",
    ) as writer:
        attempts_df.to_excel(
            writer,
            sheet_name="Attempts",
            index=False,
        )

        question_details_df.to_excel(
            writer,
            sheet_name="Question Details",
            index=False,
        )

    workbook = load_workbook(RESULTS_FILE)

    style_excel_sheet(
        workbook["Attempts"],
        "AttemptsTable",
    )

    style_excel_sheet(
        workbook["Question Details"],
        "QuestionDetailsTable",
    )

    workbook.save(RESULTS_FILE)


def choose_option(option):
    if st.session_state.selected_option is not None:
        return

    question_index = st.session_state.question_index
    question = QUESTIONS[question_index]

    st.session_state.selected_option = option

    st.session_state.answers.append(
        {
            "question": question_index + 1,
            "selected_option": option,
            "correct_option": question["correct_option"],
            "is_correct": option == question["correct_option"],
            "distance_option_a": question["distance_option_a"],
            "distance_option_b": question["distance_option_b"],
        }
    )


def next_question():
    st.session_state.question_index += 1
    st.session_state.selected_option = None

    if st.session_state.question_index >= len(QUESTIONS):
        st.session_state.quiz_finished = True


def restart_quiz():
    st.session_state.screen = "study"
    st.session_state.question_index = 0
    st.session_state.answers = []
    st.session_state.selected_option = None
    st.session_state.quiz_finished = False


def show_image(image_path):
    """Show a fixed-height image for the quiz question page."""
    if Path(image_path).exists():
        with st.container():
            st.markdown(
                '<div class="quiz-image">',
                unsafe_allow_html=True,
            )
            st.image(
                image_path,
                width="stretch",
            )
            st.markdown(
                "</div>",
                unsafe_allow_html=True,
            )
    else:
        st.error(f"Image not found: `{image_path}`")


def show_example_image(image_path):
    """Show the full-height instructional diagram."""
    if Path(image_path).exists():
        with st.container():
            st.markdown(
                '<div class="example-diagram">',
                unsafe_allow_html=True,
            )
            st.image(
                image_path,
                width="stretch",
            )
            st.markdown(
                "</div>",
                unsafe_allow_html=True,
            )
    else:
        st.error(f"Example diagram not found: `{image_path}`")


def show_question_label(question_number):
    st.markdown(
        f"""
        <div style="
            height: {IMAGE_HEIGHT_PX}px;
            display: flex;
            align-items: center;
            justify-content: center;
            font-size: 1.25rem;
            font-weight: 700;
        ">
            Q{question_number}
        </div>
        """,
        unsafe_allow_html=True,
    )


def show_option_button(option, question_index):
    return st.button(
        option,
        key=f"choose_{option.lower()}_{question_index}",
        width="stretch",
        disabled=st.session_state.selected_option is not None,
    )


def show_study_slide():
    st.title("Visual Similarity of Distributions")

    st.write(
        """
        This study investigates how people visually compare data
        distributions shown as scatter plots.
        """
    )

    st.write(
        """
        We are interested in whether people can identify which of two
        distributions appears more similar to a reference distribution.
        """
    )

    st.subheader("What is Wasserstein distance?")

    st.write(
        """
        Wasserstein distance is a mathematical measure of the difference
        between two distributions.
        """
    )

    st.write(
        """
        It can be understood as the minimum effort needed to move the
        points of one distribution so that they match another distribution.
        """
    )

    left_column, right_column = st.columns(2)

    with left_column:
        st.info(
            """
            **Smaller Wasserstein distance**

            The distributions are more similar because less movement is
            needed to transform one into the other.
            """
        )

    with right_column:
        st.warning(
            """
            **Larger Wasserstein distance**

            The distributions are less similar because more movement is
            needed to transform one into the other.
            """
        )

    st.subheader("Example")

    if Path(EXAMPLE_DIAGRAM_PATH).exists():
        # This uses a separate function so the introductory diagram
        # displays at its natural height and Option B remains visible.
        show_example_image(EXAMPLE_DIAGRAM_PATH)

        example_distance_a_column, example_distance_b_column = st.columns(2)

        with example_distance_a_column:
            st.metric(
                label="Wasserstein Distance: Option A to Q",
                value=f"{EXAMPLE_DISTANCE_A_TO_Q:.4f}",
            )

        with example_distance_b_column:
            st.metric(
                label="Wasserstein Distance: Option B to Q",
                value=f"{EXAMPLE_DISTANCE_B_TO_Q:.4f}",
            )

        st.markdown(
            """
            <div class="slide-note">
                In this example, <strong>Option A</strong> and
                <strong>Reference Q</strong> are both compact and located in
                similar positions. Option A therefore has a smaller
                Wasserstein distance to Q.
                <strong>Option B</strong> is wider and farther from Q. Its
                supports would need to move farther to match Q, so it has a
                larger Wasserstein distance.
            </div>
            """,
            unsafe_allow_html=True,
        )

    else:
        st.markdown(
            """
            <div class="diagram-placeholder">
                <strong>Diagram placeholder</strong><br><br>
                Add your example scatter-plot diagram here:<br>
                <code>images/example/wasserstein_example.png</code><br><br>
                Suggested layout:<br>
                Option A, Reference Q, Option B<br>
                with a smaller W(A, Q) and a larger W(B, Q).
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.button(
        "Next",
        type="primary",
        width="stretch",
        on_click=go_to_instructions,
    )


def show_instructions_slide():
    st.title("Instructions")

    st.write("The study contains **20 questions**.")

    st.write(
        """
        Each question contains three scatter plots arranged vertically:
        """
    )

    st.markdown(
        """
        ```text
        [A]  Option A distribution

        [Q]  Reference distribution

        [B]  Option B distribution
        ```
        """
    )

    st.subheader("Your task")

    st.write(
        """
        For each question, choose the option, **A** or **B**, that looks
        more similar to the reference distribution **Q**.
        """
    )

    st.subheader("What you may consider")

    st.markdown(
        """
        - Overall position or location
        - Shape and orientation
        - Spread or concentration
        - Clusters or separated groups
        - Outliers
        """
    )

    st.subheader("After selecting an answer")

    st.markdown(
        """
        1. You will see whether your answer is correct.
        2. You will see the Wasserstein distance from Option A to Q.
        3. You will see the Wasserstein distance from Option B to Q.
        4. Click **Next Question** to continue.
        """
    )

    st.markdown(
        """
        <div class="slide-note">
            <strong>No calculation is required.</strong><br>
            Please choose the option that looks more similar to Q.
        </div>
        """,
        unsafe_allow_html=True,
    )

    back_column, start_column = st.columns(2)

    with back_column:
        st.button(
            "Back",
            width="stretch",
            on_click=lambda: st.session_state.update({"screen": "study"}),
        )

    with start_column:
        st.button(
            "Start Study",
            type="primary",
            width="stretch",
            on_click=start_quiz,
        )


def show_quiz():
    index = st.session_state.question_index
    question = QUESTIONS[index]

    st.progress((index + 1) / len(QUESTIONS))
    st.subheader(f"Question {index + 1} of {len(QUESTIONS)}")
    st.write(QUESTION_TEXT)

    # Row 1: Option A
    option_a_button, option_a_image = st.columns(
        [1, 8],
        vertical_alignment="center",
    )

    with option_a_button:
        option_a_selected = show_option_button("A", index)

    with option_a_image:
        show_image(question["option_a_image"])

    if option_a_selected:
        choose_option("A")
        st.rerun()

    # Row 2: Reference distribution Q
    target_label, target_image = st.columns(
        [1, 8],
        vertical_alignment="center",
    )

    with target_label:
        show_question_label(index + 1)

    with target_image:
        show_image(question["target_image"])

    # Row 3: Option B
    option_b_button, option_b_image = st.columns(
        [1, 8],
        vertical_alignment="center",
    )

    with option_b_button:
        option_b_selected = show_option_button("B", index)

    with option_b_image:
        show_image(question["option_b_image"])

    if option_b_selected:
        choose_option("B")
        st.rerun()

    if st.session_state.selected_option is not None:
        selected = st.session_state.selected_option
        correct = question["correct_option"]

        if selected == correct:
            st.success("Correct.")
        else:
            st.error(f"Incorrect. The correct answer is Option {correct}.")

        distance_a_column, distance_b_column = st.columns(2)

        with distance_a_column:
            st.metric(
                label="Wasserstein Distance: Option A to Q",
                value=f"{question['distance_option_a']:.4f}",
            )

        with distance_b_column:
            st.metric(
                label="Wasserstein Distance: Option B to Q",
                value=f"{question['distance_option_b']:.4f}",
            )

    is_last_question = index == len(QUESTIONS) - 1
    next_button_text = "Finish Study" if is_last_question else "Next Question"

    if st.button(
        next_button_text,
        type="primary",
        width="stretch",
        disabled=st.session_state.selected_option is None,
    ):
        if is_last_question:
            record_result()
            st.session_state.quiz_finished = True
        else:
            next_question()

        st.rerun()

def show_results():
    score = sum(answer["is_correct"] for answer in st.session_state.answers)
    total = len(QUESTIONS)
    percentage = 100 * score / total

    st.header("Study Complete")

    score_column, percentage_column = st.columns(2)

    with score_column:
        st.metric(
            label="Final Score",
            value=f"{score} / {total}",
        )

    with percentage_column:
        st.metric(
            label="Percentage",
            value=f"{percentage:.1f}%",
        )

    results_rows = [
        {
            "Question": f"Q{answer['question']}",
            "Your Answer": f"Option {answer['selected_option']}",
            "Correct Answer": f"Option {answer['correct_option']}",
            "Result": "Correct" if answer["is_correct"] else "Incorrect",
            "W(A, Q)": answer["distance_option_a"],
            "W(B, Q)": answer["distance_option_b"],
        }
        for answer in st.session_state.answers
    ]

    results_df = pd.DataFrame(results_rows)

    st.subheader("Answer Summary")

    st.dataframe(
        results_df,
        width="stretch",
        hide_index=True,
        column_config={
            "Question": st.column_config.TextColumn(
                "Question",
                width="small",
            ),
            "Your Answer": st.column_config.TextColumn(
                "Your Answer",
                width="medium",
            ),
            "Correct Answer": st.column_config.TextColumn(
                "Correct Answer",
                width="medium",
            ),
            "Result": st.column_config.TextColumn(
                "Result",
                width="medium",
            ),
            "W(A, Q)": st.column_config.NumberColumn(
                "W(A, Q)",
                format="%.4f",
                width="medium",
            ),
            "W(B, Q)": st.column_config.NumberColumn(
                "W(B, Q)",
                format="%.4f",
                width="medium",
            ),
        },
    )

    # Convert the table displayed above into a downloadable CSV.
    csv_data = results_df.to_csv(
        index=False,
    ).encode("utf-8")

    st.download_button(
        label="Download Results",
        data=csv_data,
        file_name="wasserstein_similarity_results.csv",
        mime="text/csv",
        type="primary",
        width="stretch",
    )

    st.button(
        "Restart Study",
        width="stretch",
        on_click=restart_quiz,
    )

initialize_state()

if st.session_state.screen == "study":
    show_study_slide()

elif st.session_state.screen == "instructions":
    show_instructions_slide()

elif st.session_state.quiz_finished:
    show_results()

else:
    show_quiz()