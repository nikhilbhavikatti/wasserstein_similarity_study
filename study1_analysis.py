from pathlib import Path
import re

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


# ---------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------

PARTICIPANT_RESULTS_DIR = Path("participant_results")
ANSWER_KEY_FILE = Path("answer_key.csv")
OUTPUT_DIR = Path("output")

QUESTION_COLUMN = "Question"
PARTICIPANT_ANSWER_COLUMN = "Your Answer"
ANSWER_KEY_COLUMN = "Correct Answer"
DISTANCE_A_COLUMN = "W(A, Q)"
DISTANCE_B_COLUMN = "W(B, Q)"

LIGHT_GREEN = "C6EFCE"
LIGHT_RED = "FCE4D6"
HEADER_BLUE = "1F4E78"


# ---------------------------------------------------------------------
# Utility functions
# ---------------------------------------------------------------------

def question_number(question_id):
    """Convert Q12 or 12 into integer 12 for reliable question sorting."""
    match = re.fullmatch(r"\s*Q?(\d+)\s*", str(question_id))

    if match is None:
        raise ValueError(
            f"Invalid question label '{question_id}'. "
            "Expected labels such as Q1, Q2, ..., Q20."
        )

    return int(match.group(1))


def normalise_question(value):
    """Convert question values to a common format: Q1, Q2, ..."""
    return f"Q{question_number(value)}"


def normalise_option(value):
    """Convert A, B, Option A, or Option B into A or B."""
    text = str(value).strip().upper()

    if text in {"A", "OPTION A"}:
        return "A"

    if text in {"B", "OPTION B"}:
        return "B"

    raise ValueError(
        f"Invalid option value '{value}'. "
        "Expected A, B, Option A, or Option B."
    )


def optional_normalise_option(value):
    """Normalise an option value when present."""
    if pd.isna(value):
        return np.nan

    return normalise_option(value)


def require_columns(dataframe, required_columns, file_label):
    """Raise an informative error if a file is missing expected columns."""
    missing_columns = [
        column
        for column in required_columns
        if column not in dataframe.columns
    ]

    if missing_columns:
        raise ValueError(
            f"{file_label} is missing required column(s): {missing_columns}\n"
            f"Columns found: {list(dataframe.columns)}"
        )


def save_dataframe(dataframe, filename):
    """Save a dataframe as CSV without the pandas index."""
    dataframe.to_csv(
        OUTPUT_DIR / filename,
        index=False,
    )


def add_value_labels(axis, bars, value_format="{:.1f}%"):
    """Write values above bar-chart bars."""
    for bar in bars:
        height = bar.get_height()

        axis.annotate(
            value_format.format(height),
            xy=(
                bar.get_x() + bar.get_width() / 2,
                height,
            ),
            xytext=(0, 4),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=10,
        )


def style_excel_sheet(worksheet, table_name):
    """Apply readable formatting and an Excel table."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor=HEADER_BLUE,
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32

    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter

        maximum_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )

        worksheet.column_dimensions[column_letter].width = min(
            max(maximum_length + 3, 12),
            28,
        )

    if worksheet.max_row >= 2:
        table = Table(
            displayName=table_name,
            ref=worksheet.dimensions,
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = table_style
        worksheet.add_table(table)


def export_response_matrix_excel(
    response_matrix_with_correct,
    participant_columns,
):
    """
    Export the response matrix to Excel.

    A participant cell is light green when that participant selected
    the correct option for the corresponding question. Incorrect cells
    remain unfilled.
    """
    output_path = OUTPUT_DIR / "response_matrix.xlsx"

    response_matrix_with_correct.to_excel(
        output_path,
        sheet_name="Response Matrix",
        index=False,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Response Matrix"]

    style_excel_sheet(
        worksheet,
        "ResponseMatrixTable",
    )

    correct_fill = PatternFill(
        fill_type="solid",
        fgColor=LIGHT_GREEN,
    )

    header_lookup = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    correct_answer_column = header_lookup["Correct Option"]

    for row_index in range(2, worksheet.max_row + 1):
        correct_option = worksheet.cell(
            row=row_index,
            column=correct_answer_column,
        ).value

        for participant in participant_columns:
            participant_column = header_lookup[participant]

            participant_answer = worksheet.cell(
                row=row_index,
                column=participant_column,
            ).value

            if participant_answer == correct_option:
                worksheet.cell(
                    row=row_index,
                    column=participant_column,
                ).fill = correct_fill

    workbook.save(output_path)


# ---------------------------------------------------------------------
# Input loading
# ---------------------------------------------------------------------

def load_answer_key():
    """
    Read answer_key.csv.

    Required:
    - Question
    - Correct Answer

    Optional:
    - W(A, Q)
    - W(B, Q)
    """
    if not ANSWER_KEY_FILE.exists():
        raise FileNotFoundError(
            f"Answer key not found: {ANSWER_KEY_FILE}"
        )

    answer_key = pd.read_csv(ANSWER_KEY_FILE)

    require_columns(
        answer_key,
        [QUESTION_COLUMN, ANSWER_KEY_COLUMN],
        f"Answer key '{ANSWER_KEY_FILE}'",
    )

    answer_key = answer_key.copy()

    answer_key[QUESTION_COLUMN] = answer_key[QUESTION_COLUMN].apply(
        normalise_question
    )

    answer_key["correct_option"] = answer_key[
        ANSWER_KEY_COLUMN
    ].apply(
        normalise_option
    )

    if answer_key[QUESTION_COLUMN].duplicated().any():
        duplicate_questions = answer_key.loc[
            answer_key[QUESTION_COLUMN].duplicated(),
            QUESTION_COLUMN,
        ].tolist()

        raise ValueError(
            f"Answer key contains duplicate questions: "
            f"{duplicate_questions}"
        )

    if DISTANCE_A_COLUMN in answer_key.columns:
        answer_key["distance_a"] = pd.to_numeric(
            answer_key[DISTANCE_A_COLUMN],
            errors="raise",
        )

    if DISTANCE_B_COLUMN in answer_key.columns:
        answer_key["distance_b"] = pd.to_numeric(
            answer_key[DISTANCE_B_COLUMN],
            errors="raise",
        )

    keep_columns = [
        QUESTION_COLUMN,
        "correct_option",
        "distance_a",
        "distance_b",
    ]

    keep_columns = [
        column
        for column in keep_columns
        if column in answer_key.columns
    ]

    return answer_key[keep_columns]


def load_participant_results():
    """
    Load any number of participant CSV files.

    Required participant columns:
    - Question
    - Your Answer

    Optional participant columns:
    - Correct Answer
    - W(A, Q)
    - W(B, Q)
    """
    if not PARTICIPANT_RESULTS_DIR.exists():
        raise FileNotFoundError(
            f"Participant-results folder not found: "
            f"{PARTICIPANT_RESULTS_DIR}"
        )

    participant_files = sorted(
        PARTICIPANT_RESULTS_DIR.glob("*.csv")
    )

    if not participant_files:
        raise FileNotFoundError(
            f"No participant CSV files found in: "
            f"{PARTICIPANT_RESULTS_DIR}"
        )

    all_participant_frames = []

    for file_path in participant_files:
        participant_id = file_path.stem
        participant_df = pd.read_csv(file_path)

        require_columns(
            participant_df,
            [QUESTION_COLUMN, PARTICIPANT_ANSWER_COLUMN],
            f"Participant file '{file_path.name}'",
        )

        participant_df = participant_df.copy()

        participant_df[QUESTION_COLUMN] = participant_df[
            QUESTION_COLUMN
        ].apply(
            normalise_question
        )

        participant_df["selected_option"] = participant_df[
            PARTICIPANT_ANSWER_COLUMN
        ].apply(
            normalise_option
        )

        if participant_df[QUESTION_COLUMN].duplicated().any():
            duplicate_questions = participant_df.loc[
                participant_df[QUESTION_COLUMN].duplicated(),
                QUESTION_COLUMN,
            ].tolist()

            raise ValueError(
                f"Participant file '{file_path.name}' contains duplicate "
                f"questions: {duplicate_questions}"
            )

        participant_df["participant"] = participant_id

        if ANSWER_KEY_COLUMN in participant_df.columns:
            participant_df["participant_file_correct_option"] = (
                participant_df[ANSWER_KEY_COLUMN].apply(
                    optional_normalise_option
                )
            )

        if DISTANCE_A_COLUMN in participant_df.columns:
            participant_df["participant_file_distance_a"] = pd.to_numeric(
                participant_df[DISTANCE_A_COLUMN],
                errors="raise",
            )

        if DISTANCE_B_COLUMN in participant_df.columns:
            participant_df["participant_file_distance_b"] = pd.to_numeric(
                participant_df[DISTANCE_B_COLUMN],
                errors="raise",
            )

        keep_columns = [
            QUESTION_COLUMN,
            "participant",
            "selected_option",
            "participant_file_correct_option",
            "participant_file_distance_a",
            "participant_file_distance_b",
        ]

        keep_columns = [
            column
            for column in keep_columns
            if column in participant_df.columns
        ]

        all_participant_frames.append(
            participant_df[keep_columns]
        )

    return pd.concat(
        all_participant_frames,
        ignore_index=True,
    )


def add_or_validate_distance_information(answer_key, responses):
    """
    Use distances from the answer key where available.

    If they are absent from answer_key.csv, derive the distance values
    from the participant files and verify that values agree across files.
    """
    answer_key_has_distances = (
        "distance_a" in answer_key.columns
        and "distance_b" in answer_key.columns
    )

    if answer_key_has_distances:
        return answer_key

    required_columns = [
        "participant_file_distance_a",
        "participant_file_distance_b",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in responses.columns
    ]

    if missing_columns:
        raise ValueError(
            "Distances are missing from answer_key.csv and participant "
            f"files. Missing columns: {missing_columns}"
        )

    distance_consistency = (
        responses.groupby(QUESTION_COLUMN, as_index=False)
        .agg(
            unique_a=("participant_file_distance_a", "nunique"),
            unique_b=("participant_file_distance_b", "nunique"),
        )
    )

    inconsistent = distance_consistency[
        (distance_consistency["unique_a"] > 1)
        | (distance_consistency["unique_b"] > 1)
    ]

    if not inconsistent.empty:
        raise ValueError(
            "Distance values differ across participant files for:\n"
            f"{inconsistent.to_string(index=False)}"
        )

    distance_by_question = (
        responses.groupby(QUESTION_COLUMN, as_index=False)
        .agg(
            distance_a=("participant_file_distance_a", "first"),
            distance_b=("participant_file_distance_b", "first"),
        )
    )

    return answer_key.merge(
        distance_by_question,
        on=QUESTION_COLUMN,
        how="left",
        validate="one_to_one",
    )


def validate_participant_file_answers(responses, answer_key):
    """
    If participant CSVs include Correct Answer, check those values against
    answer_key.csv. The external answer key remains authoritative.
    """
    if "participant_file_correct_option" not in responses.columns:
        return

    expected_answers = answer_key[
        [QUESTION_COLUMN, "correct_option"]
    ].rename(
        columns={
            "correct_option": "answer_key_correct_option",
        }
    )

    comparison = responses.merge(
        expected_answers,
        on=QUESTION_COLUMN,
        how="left",
        validate="many_to_one",
    )

    available_answers = comparison.dropna(
        subset=["participant_file_correct_option"]
    )

    mismatches = available_answers[
        available_answers["participant_file_correct_option"]
        != available_answers["answer_key_correct_option"]
    ]

    if not mismatches.empty:
        raise ValueError(
            "Correct Answer values in participant result files do not "
            "match answer_key.csv:\n"
            f"{mismatches.to_string(index=False)}"
        )


# ---------------------------------------------------------------------
# Plot functions
# ---------------------------------------------------------------------

def create_participant_accuracy_plot(participant_summary):
    """Create participant accuracy bar plot with y-axis up to 110%."""

    figure, axis = plt.subplots(
        figsize=(max(8, len(participant_summary) * 1.3), 5.2),
        layout="constrained",
    )

    bars = axis.bar(
        participant_summary["Participant"],
        participant_summary["Accuracy (%)"],
        color="tab:blue",
    )

    axis.set_ylim(0, 110)
    axis.set_ylabel("Accuracy (%)")
    axis.set_xlabel("Participant")
    axis.set_title(
        "Participant Accuracy",
        fontweight="bold",
    )

    axis.grid(
        axis="y",
        linestyle="--",
        alpha=0.4,
    )

    add_value_labels(axis, bars)

    figure.savefig(
        OUTPUT_DIR / "participant_accuracy.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


def create_question_accuracy_plot(question_summary):
    """Create question accuracy bar plot with y-axis up to 110%."""

    figure, axis = plt.subplots(
        figsize=(14, 5.2),
        layout="constrained",
    )

    colors = [
        "tab:green"
        if accuracy >= 75
        else "tab:orange"
        if accuracy >= 50
        else "tab:red"
        for accuracy in question_summary["Accuracy (%)"]
    ]

    bars = axis.bar(
        question_summary["Question"],
        question_summary["Accuracy (%)"],
        color=colors,
    )

    axis.set_ylim(0, 110)
    axis.set_ylabel("Participant Accuracy (%)")
    axis.set_xlabel("Question")
    axis.set_title(
        "Question Accuracy Across Participants",
        fontweight="bold",
    )

    axis.grid(
        axis="y",
        linestyle="--",
        alpha=0.4,
    )

    add_value_labels(axis, bars)

    figure.savefig(
        OUTPUT_DIR / "question_accuracy.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


def create_response_heatmap(
    response_matrix,
    correct_options_by_question,
):
    """
    Create a response heatmap.

    Light green means the participant selected the correct option.
    White means the participant selected the incorrect option.
    The letters A and B are written inside every cell.
    """
    participants = list(response_matrix.columns)
    questions = list(response_matrix.index)

    correct_mask = np.zeros(
        shape=response_matrix.shape,
        dtype=float,
    )

    for question_index, question in enumerate(questions):
        correct_option = correct_options_by_question.loc[question]

        for participant_index, participant in enumerate(participants):
            selected_option = response_matrix.loc[
                question,
                participant,
            ]

            correct_mask[
                question_index,
                participant_index,
            ] = int(selected_option == correct_option)

    figure_width = max(8, len(participants) * 1.25 + 4)
    figure_height = max(8, len(questions) * 0.42 + 2)

    figure, axis = plt.subplots(
        figsize=(figure_width, figure_height),
        layout="constrained",
    )

    cmap = plt.matplotlib.colors.ListedColormap(
        ["#FFFFFF", "#C6EFCE"]
    )

    axis.imshow(
        correct_mask,
        cmap=cmap,
        vmin=0,
        vmax=1,
        aspect="auto",
    )

    axis.set_xticks(range(len(participants)))
    axis.set_xticklabels(
        participants,
        rotation=35,
        ha="right",
    )

    axis.set_yticks(range(len(questions)))
    axis.set_yticklabels(questions)

    axis.set_xlabel("Participant")
    axis.set_ylabel("Question")
    axis.set_title(
        "Participant Response Matrix",
        fontweight="bold",
    )

    for question_index, question in enumerate(questions):
        for participant_index, participant in enumerate(participants):
            selected_option = response_matrix.loc[
                question,
                participant,
            ]

            axis.text(
                participant_index,
                question_index,
                selected_option,
                ha="center",
                va="center",
                fontsize=10,
                fontweight="bold",
                color="black",
            )

    legend_handles = [
        plt.matplotlib.patches.Patch(
            facecolor="#C6EFCE",
            edgecolor="#9BBB59",
            label="Correct selection",
        ),
        plt.matplotlib.patches.Patch(
            facecolor="#FFFFFF",
            edgecolor="#BFBFBF",
            label="Incorrect selection",
        ),
    ]

    axis.legend(
        handles=legend_handles,
        loc="upper center",
        bbox_to_anchor=(0.5, -0.07),
        ncol=2,
        frameon=False,
    )

    figure.savefig(
        OUTPUT_DIR / "response_heatmap.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


def create_pairwise_agreement_heatmap(pairwise_agreement):
    """Create pairwise participant agreement heatmap."""

    values = pairwise_agreement.to_numpy(dtype=float)
    participants = list(pairwise_agreement.columns)

    figure_width = max(7, len(participants) * 1.2 + 3)
    figure_height = max(6, len(participants) * 1.0 + 2)

    figure, axis = plt.subplots(
        figsize=(figure_width, figure_height),
        layout="constrained",
    )

    image = axis.imshow(
        values,
        cmap="Blues",
        vmin=0,
        vmax=100,
        aspect="auto",
    )

    axis.set_xticks(range(len(participants)))
    axis.set_xticklabels(
        participants,
        rotation=35,
        ha="right",
    )

    axis.set_yticks(range(len(participants)))
    axis.set_yticklabels(participants)

    axis.set_xlabel("Participant")
    axis.set_ylabel("Participant")
    axis.set_title(
        "Pairwise Participant Agreement",
        fontweight="bold",
    )

    for row_index in range(values.shape[0]):
        for column_index in range(values.shape[1]):
            value = values[row_index, column_index]

            text_color = "white" if value >= 60 else "black"

            axis.text(
                column_index,
                row_index,
                f"{value:.1f}%",
                ha="center",
                va="center",
                fontsize=10,
                color=text_color,
            )

    colorbar = figure.colorbar(
        image,
        ax=axis,
        shrink=0.85,
    )

    colorbar.ax.set_ylabel(
        "Agreement (%)",
        rotation=270,
        labelpad=16,
    )

    figure.savefig(
        OUTPUT_DIR / "pairwise_agreement_heatmap.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


# ---------------------------------------------------------------------
# Main analysis
# ---------------------------------------------------------------------

def main():
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    answer_key = load_answer_key()
    responses = load_participant_results()

    answer_key = add_or_validate_distance_information(
        answer_key,
        responses,
    )

    validate_participant_file_answers(
        responses,
        answer_key,
    )

    expected_questions = set(answer_key[QUESTION_COLUMN])
    participant_ids = sorted(responses["participant"].unique())

    # Check that every participant has exactly the same question set
    # as the answer key. This works for any number of participants.
    for participant in participant_ids:
        answered_questions = set(
            responses.loc[
                responses["participant"] == participant,
                QUESTION_COLUMN,
            ]
        )

        missing_questions = sorted(
            expected_questions - answered_questions,
            key=question_number,
        )

        extra_questions = sorted(
            answered_questions - expected_questions,
            key=question_number,
        )

        if missing_questions or extra_questions:
            raise ValueError(
                f"Question mismatch for participant '{participant}'.\n"
                f"Missing: {missing_questions}\n"
                f"Unexpected: {extra_questions}"
            )

    merged = responses.merge(
        answer_key,
        on=QUESTION_COLUMN,
        how="left",
        validate="many_to_one",
    )

    if merged["correct_option"].isna().any():
        missing_key_questions = merged.loc[
            merged["correct_option"].isna(),
            QUESTION_COLUMN,
        ].unique()

        raise ValueError(
            "Questions missing from answer_key.csv: "
            f"{list(missing_key_questions)}"
        )

    if merged["distance_a"].isna().any() or merged["distance_b"].isna().any():
        missing_distance_questions = merged.loc[
            merged["distance_a"].isna() | merged["distance_b"].isna(),
            QUESTION_COLUMN,
        ].unique()

        raise ValueError(
            "Missing Wasserstein distances for questions: "
            f"{list(missing_distance_questions)}"
        )

    merged["is_correct"] = (
        merged["selected_option"] == merged["correct_option"]
    )

    merged["distance_gap"] = np.abs(
        merged["distance_a"] - merged["distance_b"]
    )

    merged["Question Number"] = merged[QUESTION_COLUMN].apply(
        question_number
    )

    merged = merged.sort_values(
        ["participant", "Question Number"],
    ).reset_index(drop=True)

    # -----------------------------------------------------------------
    # 1. All merged responses
    # -----------------------------------------------------------------

    all_responses_export = merged[
        [
            "participant",
            QUESTION_COLUMN,
            "selected_option",
            "correct_option",
            "is_correct",
            "distance_a",
            "distance_b",
            "distance_gap",
        ]
    ].rename(
        columns={
            "participant": "Participant",
            QUESTION_COLUMN: "Question",
            "selected_option": "Selected Option",
            "correct_option": "Correct Option",
            "is_correct": "Correct",
            "distance_a": "W(A, Q)",
            "distance_b": "W(B, Q)",
            "distance_gap": "|W(A, Q) - W(B, Q)|",
        }
    )

    save_dataframe(
        all_responses_export,
        "all_responses_merged.csv",
    )

    # -----------------------------------------------------------------
    # 2. Participant accuracy summary
    # Used only to produce participant_accuracy.png.
    # It is not saved as a separate CSV/XLSX file.
    # -----------------------------------------------------------------

    participant_summary = (
        merged.groupby("participant", as_index=False)
        .agg(
            Correct=("is_correct", "sum"),
            Total=("is_correct", "size"),
        )
    )

    participant_summary["Accuracy (%)"] = (
        100
        * participant_summary["Correct"]
        / participant_summary["Total"]
    )

    participant_summary = participant_summary.rename(
        columns={
            "participant": "Participant",
        }
    )

    participant_summary["Accuracy (%)"] = participant_summary[
        "Accuracy (%)"
    ].round(2)

    # -----------------------------------------------------------------
    # 3. Question summary
    # -----------------------------------------------------------------

    response_counts = (
        merged.groupby(QUESTION_COLUMN)["selected_option"]
        .value_counts()
        .unstack(fill_value=0)
    )

    if "A" not in response_counts.columns:
        response_counts["A"] = 0

    if "B" not in response_counts.columns:
        response_counts["B"] = 0

    response_counts = response_counts[
        ["A", "B"]
    ].reset_index()

    response_counts = response_counts.rename(
        columns={
            "A": "Selected A",
            "B": "Selected B",
        }
    )

    question_summary = (
        merged.groupby(
            [
                QUESTION_COLUMN,
                "correct_option",
                "distance_a",
                "distance_b",
                "distance_gap",
            ],
            as_index=False,
        )
        .agg(
            Correct=("is_correct", "sum"),
            Participants=("is_correct", "size"),
        )
    )

    question_summary["Accuracy (%)"] = (
        100
        * question_summary["Correct"]
        / question_summary["Participants"]
    )

    question_summary = question_summary.merge(
        response_counts,
        on=QUESTION_COLUMN,
        how="left",
        validate="one_to_one",
    )

    question_summary["Majority Choice"] = np.select(
        [
            question_summary["Selected A"] > question_summary["Selected B"],
            question_summary["Selected B"] > question_summary["Selected A"],
        ],
        [
            "A",
            "B",
        ],
        default="Tie",
    )

    question_summary["Question Number"] = question_summary[
        QUESTION_COLUMN
    ].apply(question_number)

    question_summary = question_summary.sort_values(
        "Question Number"
    ).reset_index(drop=True)

    question_summary_export = question_summary[
        [
            QUESTION_COLUMN,
            "correct_option",
            "distance_a",
            "distance_b",
            "distance_gap",
            "Correct",
            "Participants",
            "Accuracy (%)",
            "Selected A",
            "Selected B",
            "Majority Choice",
        ]
    ].rename(
        columns={
            QUESTION_COLUMN: "Question",
            "correct_option": "Correct Option",
            "distance_a": "W(A, Q)",
            "distance_b": "W(B, Q)",
            "distance_gap": "Distance Gap |W(A,Q) - W(B,Q)|",
            "Correct": "Correct Participants",
        }
    )

    for column in [
        "W(A, Q)",
        "W(B, Q)",
        "Distance Gap |W(A,Q) - W(B,Q)|",
        "Accuracy (%)",
    ]:
        question_summary_export[column] = question_summary_export[
            column
        ].round(4)

    save_dataframe(
        question_summary_export,
        "question_summary.csv",
    )

    # -----------------------------------------------------------------
    # 4. Response matrix
    # -----------------------------------------------------------------

    response_matrix = merged.pivot(
        index=QUESTION_COLUMN,
        columns="participant",
        values="selected_option",
    )

    response_matrix = response_matrix.reindex(
        sorted(response_matrix.index, key=question_number)
    )

    response_matrix = response_matrix[
        sorted(response_matrix.columns)
    ]

    correct_options_by_question = (
        answer_key.set_index(QUESTION_COLUMN)["correct_option"]
        .reindex(response_matrix.index)
    )

    response_matrix_with_correct = response_matrix.copy()
    response_matrix_with_correct.insert(
        0,
        "Correct Option",
        correct_options_by_question,
    )

    response_matrix_export = (
        response_matrix_with_correct.reset_index()
        .rename(
            columns={
                QUESTION_COLUMN: "Question",
            }
        )
    )

    save_dataframe(
        response_matrix_export,
        "response_matrix.csv",
    )

    export_response_matrix_excel(
        response_matrix_with_correct=response_matrix_export,
        participant_columns=list(response_matrix.columns),
    )

    # -----------------------------------------------------------------
    # 5. Pairwise agreement
    # -----------------------------------------------------------------

    participants = list(response_matrix.columns)
    
    # Safety check: every participant filename must be unique because it
    # becomes a column name in the response matrix.
    if len(participants) != len(set(participants)):
        raise ValueError(
            "Duplicate participant IDs were found. Ensure each participant "
            "CSV file has a unique filename, for example:\n"
            "participant_01.csv\n"
            "participant_02.csv\n"
            "participant_03.csv"
        )
    
    num_participants = len(participants)
    
    pairwise_values = np.full(
        (num_participants, num_participants),
        np.nan,
        dtype=float,
    )
    
    for i, participant_a in enumerate(participants):
        for j, participant_b in enumerate(participants):
            responses_a = response_matrix[participant_a]
            responses_b = response_matrix[participant_b]
    
            valid_mask = responses_a.notna() & responses_b.notna()
    
            if valid_mask.sum() > 0:
                pairwise_values[i, j] = 100 * (
                    responses_a[valid_mask].to_numpy()
                    == responses_b[valid_mask].to_numpy()
                ).mean()
    
    pairwise_agreement = pd.DataFrame(
        pairwise_values,
        index=participants,
        columns=participants,
    )
    
    pairwise_agreement.index.name = "Participant"
    pairwise_agreement.columns.name = "Participant"
    
    pairwise_agreement_export = pairwise_agreement.reset_index()
    
    save_dataframe(
        pairwise_agreement_export,
        "pairwise_agreement.csv",
    )

    # -----------------------------------------------------------------
    # 6. Plots
    # -----------------------------------------------------------------

    create_participant_accuracy_plot(
        participant_summary,
    )

    create_question_accuracy_plot(
        question_summary_export,
    )

    create_response_heatmap(
        response_matrix=response_matrix,
        correct_options_by_question=correct_options_by_question,
    )

    create_pairwise_agreement_heatmap(
        pairwise_agreement,
    )

    print("\nAnalysis completed successfully.")
    print(f"Participants found: {len(participant_ids)}")
    print(f"Questions found: {len(answer_key)}")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")

    print("\nCreated files:")
    print("- all_responses_merged.csv")
    print("- question_summary.csv")
    print("- response_matrix.csv")
    print("- response_matrix.xlsx")
    print("- pairwise_agreement.csv")
    print("- participant_accuracy.png")
    print("- question_accuracy.png")
    print("- response_heatmap.png")
    print("- pairwise_agreement_heatmap.png")


if __name__ == "__main__":
    main()